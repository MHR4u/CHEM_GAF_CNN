"""
make_final_report.py
====================

Builds the FINAL RESULT REPORT of the GAF-CNN in three formats, from the files
already produced by gaf_cnn_full_output.py / make_full_report.py:

    results/CNN_Final_Report.pdf    (A4, ready to print/share)
    results/CNN_Final_Report.docx   (Word, editable)
    results/CNN_Final_Report.txt    (plain text)

The report contains: model & dataset summary, method, overall accuracy,
per-class accuracy, confusion matrices (as tables and as images), classification
reports, per-image predictions, misclassified images, an appendix with the full
probability matrix, and the list of generated files.

Usage:
    python3 make_final_report.py --results_dir ../results --data_dir ../data/gafimages \
        --model CNN_5classes__epo_4000_pat_1000_batch_16_tl.keras
"""

import argparse
import csv
import glob
import json
import os
import zipfile
from datetime import datetime, timezone

# --------------------------------------------------------------------------- #
#  data collection
# --------------------------------------------------------------------------- #

CLASSES = [
    "Benzochinone",
    "Catecolo",
    "Ferro-Ferri",
    "Idrochinone",
    "Idrochinone-Benzochinone",
]
SPLITS = ["shuffled", "concentrations_based"]


def model_info(model_path):
    info = {"file": os.path.basename(model_path), "keras_version": "-", "date_saved": "-"}
    try:
        with zipfile.ZipFile(model_path) as z:
            meta = json.loads(z.read("metadata.json").decode())
            info["keras_version"] = meta.get("keras_version", "-")
            info["date_saved"] = meta.get("date_saved", "-")
    except Exception:
        pass
    return info


def read_predictions(path):
    with open(path) as f:
        return list(csv.DictReader(f))


def dataset_inventory(data_dir):
    """Counts of images per layout / split / class."""
    rows = []
    root = os.path.dirname(os.path.abspath(data_dir))
    for path in sorted(glob.glob(os.path.join(root, "**", "*.png"), recursive=True)):
        rel = os.path.relpath(path, root)
        parts = rel.split(os.sep)
        if parts[0] == "gafimages" and len(parts) >= 4:
            rows.append(("CNN input image", parts[1], parts[2], parts[3]))
        elif parts[0] == "voltammetry" and len(parts) >= 5:
            rows.append(("raw GAF output", "voltammetry/" + parts[2], "-", parts[1]))
    counts = {}
    for kind, layout, split, cls in rows:
        key = (kind, layout, split, cls)
        counts[key] = counts.get(key, 0) + 1
    return rows, counts


def collect(results_dir, data_dir, model_path):
    summaries = json.load(open(os.path.join(results_dir, "all_splits_summary.json")))
    for s in summaries:
        s["predictions"] = read_predictions(os.path.join(results_dir, f"{s['split']}_predictions.csv"))
        with open(os.path.join(results_dir, f"{s['split']}_classification_report.txt")) as f:
            s["report_text"] = "\n".join(f.read().splitlines()[3:]).strip("\n")
        rows = s["predictions"]
        correct = [r for r in rows if r["correct"] == "True"]
        wrong = [r for r in rows if r["correct"] != "True"]
        s["avg_conf_correct"] = (sum(float(r["max_confidence"]) for r in correct) / len(correct)) if correct else 0.0
        s["avg_conf_wrong"] = (sum(float(r["max_confidence"]) for r in wrong) / len(wrong)) if wrong else 0.0
        s["wrong_rows"] = wrong
    return summaries


def fmt(v, nd=2):
    return f"{v:.{nd}f}"


# --------------------------------------------------------------------------- #
#  plain text report
# --------------------------------------------------------------------------- #

def build_txt(summaries, mi, counts, out_path, results_dir, now):
    def rule(ch="-"):
        return ch * 96

    L = []
    L.append(rule("="))
    L.append("        GAF-CNN  -  FINAL RESULT REPORT")
    L.append("        Electrochemical image classification (Gramian Angular Field + CNN)")
    L.append(rule("="))
    L.append("")
    L.append(f"Generated (UTC) : {now}")
    L.append(f"Model file      : {mi['file']}")
    L.append(f"Model parameters: {summaries[0]['model_params']:,}")
    L.append(f"Keras version   : {mi['keras_version']}   (saved {mi['date_saved']})")
    L.append(f"Input size      : 224 x 224 x 3 (GAF images, 100x100 RGB on disk)")
    L.append(f"Classes ({len(CLASSES)})     : " + ", ".join(CLASSES))
    L.append("")

    L.append(rule("="))
    L.append("1. DATASET USED")
    L.append(rule("="))
    L.append("")
    L.append(f"{'kind':<18}{'layout':<24}{'split':<8}{'class':<28}{'images':>7}")
    L.append(rule())
    for (kind, layout, split, cls), n in sorted(counts.items()):
        L.append(f"{kind:<18}{layout:<24}{split:<8}{cls:<28}{n:>7}")
    L.append(rule())
    L.append(f"{'':<58}{'TOTAL':<28}{sum(counts.values()):>7}")
    L.append("")

    L.append(rule("="))
    L.append("2. OVERALL RESULTS")
    L.append(rule("="))
    L.append("")
    L.append(f"{'test split':<24}{'images':>8}{'correct':>9}{'wrong':>7}{'accuracy %':>12}")
    L.append(rule())
    for s in summaries:
        L.append(
            f"{s['split']:<24}{s['n_test_images']:>8}{s['n_correct']:>9}{s['n_wrong']:>7}"
            f"{s['accuracy_percent']:>12.2f}"
        )
    L.append("")

    L.append(rule("="))
    L.append("3. ACCURACY PER CLASS (%)")
    L.append(rule("="))
    L.append("")
    L.append(f"{'class':<28}" + "".join(f"{s['split']:>24}" for s in summaries))
    L.append(rule())
    for c in CLASSES:
        L.append(f"{c:<28}" + "".join(f"{s['accuracy_by_class_percent'][c]:>24.2f}" for s in summaries))
    L.append("")
    L.append("Prediction confidence (average of the winning class):")
    for s in summaries:
        L.append(
            f"  {s['split']:<24} correct predictions: {s['avg_conf_correct']:.4f}"
            f"   |   wrong predictions: {s['avg_conf_wrong']:.4f}"
        )
    L.append("")

    L.append(rule("="))
    L.append("4. CONFUSION MATRICES  (rows = true class, columns = predicted class)")
    L.append(rule("="))
    for s in summaries:
        L.append("")
        L.append(f"--- {s['split']} split ---")
        L.append(f"{'':<28}" + "".join(f"{c[:12]:>14}" for c in CLASSES))
        for i, c in enumerate(CLASSES):
            L.append(f"{c:<28}" + "".join(f"{s['confusion_matrix'][i][j]:>14}" for j in range(len(CLASSES))))
    L.append("")

    L.append(rule("="))
    L.append("5. CLASSIFICATION REPORTS (precision / recall / f1-score)")
    L.append(rule("="))
    for s in summaries:
        L.append("")
        L.append(f"--- {s['split']} split ---")
        L.append(s["report_text"])
    L.append("")

    L.append(rule("="))
    L.append("6. PER IMAGE PREDICTIONS")
    L.append(rule("="))
    for s in summaries:
        L.append("")
        L.append(f"--- {s['split']} split ({s['n_test_images']} images) ---")
        L.append(
            f"{'#':>4}  {'file':<44}{'true class':<28}{'predicted class':<28}"
            f"{'confidence':>11}  {'result':>6}"
        )
        L.append(rule())
        for i, r in enumerate(s["predictions"]):
            ok = r["correct"] == "True"
            L.append(
                f"{i:>4}  {r['file']:<44}{r['true_class']:<28}{r['predicted_class']:<28}"
                f"{float(r['max_confidence']):>11.4f}  {'OK' if ok else 'WRONG':>6}"
            )
    L.append("")

    L.append(rule("="))
    L.append("7. MISCLASSIFIED IMAGES")
    L.append(rule("="))
    L.append("")
    any_wrong = False
    for s in summaries:
        if s["wrong_rows"]:
            any_wrong = True
            L.append(f"{s['split']} split:")
            for r in s["wrong_rows"]:
                probs = " ".join(
                    f"{c}={float(r['prob_' + c]):.4f}" for c in CLASSES
                )
                L.append(
                    f"  {r['file']}  true={r['true_class']}  predicted={r['predicted_class']}"
                    f"  confidence={float(r['max_confidence']):.4f}"
                )
                L.append(f"      probabilities: {probs}")
            L.append("")
    if not any_wrong:
        L.append("none - every test image was classified correctly")
    L.append("")

    L.append(rule("="))
    L.append("8. CONCLUSION")
    L.append(rule("="))
    L.append("")
    best = max(summaries, key=lambda s: s["accuracy_percent"])
    worst = min(summaries, key=lambda s: s["accuracy_percent"])
    L.append(
        f"The trained CNN classified {sum(s['n_correct'] for s in summaries)} of "
        f"{sum(s['n_test_images'] for s in summaries)} test images correctly.\n"
        f"Best split: {best['split']} with {best['accuracy_percent']:.2f}% accuracy; "
        f"{worst['split']}: {worst['accuracy_percent']:.2f}%.\n"
        "The perfect score of the concentrations_based split (every concentration level, 0.01 mM to 2 mM)\n"
        "shows that the GAF representation keeps the information needed to recognise the substance\n"
        "independently of its concentration. The only error (shuffled split) is a Benzochinone\n"
        "image at the highest concentration (2 mM) predicted as Idrochinone with 0.60 confidence,\n"
        "the two molecules having very similar quinone/hydroquinone redox responses."
    )
    L.append("")
    L.append(rule("="))
    L.append("9. GENERATED FILES")
    L.append(rule("="))
    L.append("")
    for path in sorted(os.listdir(results_dir)):
        L.append(f"  results/{path}")
    L.append("")

    with open(out_path, "w") as f:
        f.write("\n".join(L) + "\n")


# --------------------------------------------------------------------------- #
#  Word (.docx) report
# --------------------------------------------------------------------------- #

def build_docx(summaries, mi, counts, out_path, results_dir, now):
    from docx import Document
    from docx.shared import Pt, Cm, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10)

    def heading(text, level=1):
        h = doc.add_heading(text, level=level)
        return h

    def table(headers, rows, widths=None, font_size=8.5):
        t = doc.add_table(rows=1, cols=len(headers))
        t.style = "Light Grid Accent 1"
        for i, htext in enumerate(headers):
            cell = t.rows[0].cells[i]
            cell.text = ""
            run = cell.paragraphs[0].add_run(str(htext))
            run.bold = True
            run.font.size = Pt(font_size)
        for row in rows:
            cells = t.add_row().cells
            for i, val in enumerate(row):
                cells[i].text = ""
                run = cells[i].paragraphs[0].add_run(str(val))
                run.font.size = Pt(font_size)
        if widths:
            for r in t.rows:
                for i, w in enumerate(widths):
                    r.cells[i].width = Cm(w)
        return t

    def mono(text):
        p = doc.add_paragraph()
        run = p.add_run(text)
        run.font.name = "Consolas"
        run.font.size = Pt(7.5)
        return p

    title = doc.add_heading("GAF-CNN - Final Result Report", level=0)
    sub = doc.add_paragraph("Electrochemical image classification (Gramian Angular Field + CNN)")
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph(f"Generated (UTC): {now}").alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph()

    heading("1. Model and input", 1)
    table(
        ["Item", "Value"],
        [
            ["Model file", mi["file"]],
            ["Parameters", f"{summaries[0]['model_params']:,}"],
            ["Keras version", f"{mi['keras_version']} (saved {mi['date_saved']})"],
            ["Input size", "224 x 224 x 3 (GAF image, 100x100 RGB on disk)"],
            ["Number of classes", str(len(CLASSES))],
            ["Classes", ", ".join(CLASSES)],
            ["Pre-processing", "pixel rescale 1/255, no augmentation at test time, batch 1, shuffle disabled"],
        ],
        widths=[4.5, 11.5],
    )

    heading("2. Dataset used", 1)
    rows = [[k[0], k[1], k[2], k[3], n] for k, n in sorted(counts.items())]
    table(["Kind", "Layout", "Split", "Class", "Images"], rows, widths=[3.4, 4.0, 1.8, 4.6, 1.7])
    doc.add_paragraph(f"Total images in data/: {sum(counts.values())}")

    heading("3. Overall results", 1)
    table(
        ["Test split", "Images", "Correct", "Wrong", "Accuracy %"],
        [[s["split"], s["n_test_images"], s["n_correct"], s["n_wrong"], fmt(s["accuracy_percent"])] for s in summaries],
        widths=[5.5, 2.5, 2.5, 2.5, 3.0],
    )
    doc.add_paragraph(
        f"Combined: {sum(s['n_correct'] for s in summaries)} / {sum(s['n_test_images'] for s in summaries)} "
        f"images classified correctly."
    )

    heading("4. Accuracy per class (%)", 1)
    table(
        ["Class"] + [s["split"] for s in summaries],
        [[c] + [fmt(s["accuracy_by_class_percent"][c]) for s in summaries] for c in CLASSES],
        widths=[6.5, 4.75, 4.75],
    )
    heading("4.1 Prediction confidence", 2)
    table(
        ["Test split", "Avg confidence (correct)", "Avg confidence (wrong)"],
        [[s["split"], f"{s['avg_conf_correct']:.4f}", f"{s['avg_conf_wrong']:.4f}"] for s in summaries],
        widths=[5.5, 5.25, 5.25],
    )

    heading("5. Confusion matrices", 1)
    doc.add_paragraph("Rows = true class, columns = predicted class (alphabetical order).")
    for s in summaries:
        heading(f"{s['split']} split", 2)
        table(
            [""] + [c[:12] for c in CLASSES],
            [[CLASSES[i]] + [s["confusion_matrix"][i][j] for j in range(len(CLASSES))] for i in range(len(CLASSES))],
            font_size=8,
        )
        png = os.path.join(results_dir, f"{s['split']}_confusion_matrix.png")
        if os.path.exists(png):
            doc.add_picture(png, width=Cm(11))
            doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

    heading("6. Classification reports", 1)
    for s in summaries:
        heading(f"{s['split']} split", 2)
        mono(s["report_text"])

    heading("7. Per image predictions", 1)
    for s in summaries:
        heading(f"{s['split']} split ({s['n_test_images']} images)", 2)
        table(
            ["#", "File", "True class", "Predicted class", "Confidence", "Result"],
            [
                [
                    i,
                    r["file"],
                    r["true_class"],
                    r["predicted_class"],
                    f"{float(r['max_confidence']):.4f}",
                    "OK" if r["correct"] == "True" else "WRONG",
                ]
                for i, r in enumerate(s["predictions"])
            ],
            widths=[0.9, 5.6, 3.0, 3.4, 1.8, 1.3],
            font_size=7.5,
        )

    heading("8. Misclassified images", 1)
    any_wrong = False
    for s in summaries:
        if s["wrong_rows"]:
            any_wrong = True
            for r in s["wrong_rows"]:
                doc.add_paragraph(
                    f"{s['split']}: {r['file']} - true {r['true_class']}, predicted "
                    f"{r['predicted_class']} (confidence {float(r['max_confidence']):.4f})",
                    style="List Bullet",
                )
    if not any_wrong:
        doc.add_paragraph("None - every test image was classified correctly.")

    heading("9. Conclusion", 1)
    best = max(summaries, key=lambda s: s["accuracy_percent"])
    worst = min(summaries, key=lambda s: s["accuracy_percent"])
    doc.add_paragraph(
        f"The trained CNN classified {sum(s['n_correct'] for s in summaries)} of "
        f"{sum(s['n_test_images'] for s in summaries)} test images correctly. "
        f"Best split: {best['split']} with {best['accuracy_percent']:.2f}% accuracy; "
        f"{worst['split']}: {worst['accuracy_percent']:.2f}%."
    )
    doc.add_paragraph(
        "The perfect score of the concentrations_based split (which contains every concentration level, "
        "from 0.01 mM to 2 mM) shows that the GAF representation preserves the information needed to "
        "recognise the substance independently of its concentration. The single error (shuffled split) is a "
        "Benzochinone image at the highest concentration (2 mM) predicted as Idrochinone with 0.60 confidence; "
        "the two molecules have very similar quinone / hydroquinone redox responses."
    )

    heading("10. Appendix - full probability matrix", 1)
    doc.add_paragraph("Probability assigned by the CNN to each of the 5 classes, for every test image.")
    for s in summaries:
        heading(f"{s['split']} split", 2)
        table(
            ["#", "File", "True"] + CLASSES,
            [
                [i, r["file"], r["true_class"]] + [f"{float(r['prob_' + c]):.4f}" for c in CLASSES]
                for i, r in enumerate(s["predictions"])
            ],
            widths=[0.7, 3.9, 2.0, 2.2, 2.0, 2.2, 2.0, 2.9],
            font_size=6.5,
        )

    heading("11. Generated files", 1)
    for path in sorted(os.listdir(results_dir)):
        doc.add_paragraph(f"results/{path}", style="List Bullet")

    doc.save(out_path)


# --------------------------------------------------------------------------- #
#  PDF report
# --------------------------------------------------------------------------- #

def build_pdf(summaries, mi, counts, out_path, results_dir, now):
    import matplotlib
    import reportlab
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import (
        Image,
        KeepTogether,
        PageBreak,
        Paragraph,
        SimpleDocTemplate,
        Spacer,
        Table,
        TableStyle,
    )

    font_dir = os.path.join(os.path.dirname(matplotlib.__file__), "mpl-data", "fonts", "ttf")
    pdfmetrics.registerFont(TTFont("DejaVu", os.path.join(font_dir, "DejaVuSans.ttf")))
    pdfmetrics.registerFont(TTFont("DejaVu-Bold", os.path.join(font_dir, "DejaVuSans-Bold.ttf")))
    pdfmetrics.registerFont(TTFont("DejaVuMono", os.path.join(font_dir, "DejaVuSansMono.ttf")))

    ss = getSampleStyleSheet()
    H1 = ParagraphStyle("H1", parent=ss["Heading1"], fontName="DejaVu-Bold", fontSize=13,
                        textColor=colors.HexColor("#12304d"), spaceBefore=10, spaceAfter=6)
    H2 = ParagraphStyle("H2", parent=ss["Heading2"], fontName="DejaVu-Bold", fontSize=10.5,
                        textColor=colors.HexColor("#1f4e79"), spaceBefore=8, spaceAfter=4)
    BODY = ParagraphStyle("BODY", parent=ss["BodyText"], fontName="DejaVu", fontSize=8.6, leading=12)
    SMALL = ParagraphStyle("SMALL", parent=BODY, fontSize=6.6, leading=8.4)
    MONO = ParagraphStyle("MONO", parent=BODY, fontName="DejaVuMono", fontSize=6.6, leading=8.6)
    TITLE = ParagraphStyle("TITLE", parent=ss["Title"], fontName="DejaVu-Bold", fontSize=19,
                           textColor=colors.HexColor("#12304d"), spaceAfter=4)

    def P(text, style=BODY):
        return Paragraph(str(text), style)

    def make_table(headers, rows, widths, font_size=8, align_right_from=None):
        data = [[P(h, ParagraphStyle("th", parent=SMALL, fontName="DejaVu-Bold", fontSize=font_size + 0.3))
                 for h in headers]]
        for row in rows:
            data.append([
                P(v, ParagraphStyle("td", parent=SMALL, fontSize=font_size)) for v in row
            ])
        t = Table(data, colWidths=widths, repeatRows=1, hAlign="LEFT")
        t.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#dce6f1")),
            ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#9fb3c8")),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("TOPPADDING", (0, 0), (-1, -1), 2.5),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 2.5),
            ("LEFTPADDING", (0, 0), (-1, -1), 3),
            ("RIGHTPADDING", (0, 0), (-1, -1), 3),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f4f7fa")]),
        ]))
        return t

    story = []
    story.append(P("GAF-CNN &ndash; Final Result Report", TITLE))
    story.append(P("Electrochemical image classification (Gramian Angular Field + Convolutional Neural Network)", BODY))
    story.append(P(f"Generated (UTC): {now}", BODY))
    story.append(Spacer(1, 10))

    # 1 model
    story.append(P("1. Model and input", H1))
    story.append(make_table(
        ["Item", "Value"],
        [
            ["Model file", mi["file"]],
            ["Parameters", f"{summaries[0]['model_params']:,}"],
            ["Keras version", f"{mi['keras_version']} (saved {mi['date_saved']})"],
            ["Input size", "224 x 224 x 3 (GAF image, 100 x 100 RGB on disk)"],
            ["Number of classes", str(len(CLASSES))],
            ["Classes", ", ".join(CLASSES)],
            ["Test pre-processing", "rescale 1/255, no augmentation, batch 1, shuffle disabled"],
        ],
        [4.2 * cm, 12.6 * cm], font_size=8,
    ))

    # 2 dataset
    story.append(P("2. Dataset used", H1))
    story.append(make_table(
        ["Kind", "Layout", "Split", "Class", "Images"],
        [[k[0], k[1], k[2], k[3], n] for k, n in sorted(counts.items())],
        [3.2 * cm, 4.2 * cm, 1.7 * cm, 5.2 * cm, 2.5 * cm], font_size=7.5,
    ))
    story.append(Spacer(1, 3))
    story.append(P(f"Total images in data/: {sum(counts.values())}", BODY))

    # 3 overall
    story.append(P("3. Overall results", H1))
    acc_rows = make_table(
        ["Test split", "Images", "Correct", "Wrong", "Accuracy %"],
        [[s["split"], s["n_test_images"], s["n_correct"], s["n_wrong"], fmt(s["accuracy_percent"])] for s in summaries],
        [5.0 * cm, 2.4 * cm, 2.4 * cm, 2.4 * cm, 3.0 * cm], font_size=8,
    )
    story.append(acc_rows)
    story.append(Spacer(1, 5))
    story.append(P(
        f"Combined: {sum(s['n_correct'] for s in summaries)} of {sum(s['n_test_images'] for s in summaries)} "
        f"test images classified correctly.", BODY))

    # 4 per class
    story.append(P("4. Accuracy per class (%)", H1))
    story.append(make_table(
        ["Class"] + [s["split"] for s in summaries],
        [[c] + [fmt(s["accuracy_by_class_percent"][c]) for s in summaries] for c in CLASSES],
        [6.4 * cm, 5.2 * cm, 5.2 * cm], font_size=8,
    ))
    story.append(Spacer(1, 5))
    story.append(P("Average confidence of the predicted class", H2))
    story.append(make_table(
        ["Test split", "Correct predictions", "Wrong predictions"],
        [[s["split"], f"{s['avg_conf_correct']:.4f}", f"{s['avg_conf_wrong']:.4f}"] for s in summaries],
        [6.4 * cm, 5.2 * cm, 5.2 * cm], font_size=8,
    ))

    # 5 confusion matrices
    story.append(Spacer(1, 6))
    story.append(P("5. Confusion matrices", H1))
    story.append(P("Rows = true class, columns = predicted class (alphabetical order).", BODY))
    for s in summaries:
        block = [P(f"{s['split']} split", H2)]
        block.append(make_table(
            [""] + [c[:13] for c in CLASSES],
            [[CLASSES[i]] + [s["confusion_matrix"][i][j] for j in range(len(CLASSES))] for i in range(len(CLASSES))],
            [4.0 * cm, 2.56 * cm, 2.56 * cm, 2.56 * cm, 2.56 * cm, 2.56 * cm], font_size=7.5,
        ))
        png = os.path.join(results_dir, f"{s['split']}_confusion_matrix.png")
        if os.path.exists(png):
            w = 9.2 * cm
            img = Image(png, width=w, height=w * 0.86)
            img.hAlign = "CENTER"
            block.append(Spacer(1, 3))
            block.append(img)
        block.append(Spacer(1, 12))
        story.append(KeepTogether(block))

    # 6 classification reports
    story.append(Spacer(1, 6))
    story.append(P("6. Classification reports (precision / recall / f1-score)", H1))
    for s in summaries:
        story.append(P(f"{s['split']} split", H2))
        story.append(P(s["report_text"].replace(" ", "&nbsp;").replace("\n", "<br/>"), MONO))
        story.append(Spacer(1, 6))

    # 7 per-image predictions
    story.append(Spacer(1, 6))
    story.append(P("7. Per image predictions", H1))
    for s in summaries:
        story.append(P(f"{s['split']} split &ndash; {s['n_test_images']} images", H2))
        story.append(make_table(
            ["#", "File", "True class", "Predicted class", "Conf.", "Result"],
            [
                [i, r["file"], r["true_class"], r["predicted_class"],
                 f"{float(r['max_confidence']):.4f}", "OK" if r["correct"] == "True" else "WRONG"]
                for i, r in enumerate(s["predictions"])
            ],
            [0.8 * cm, 5.6 * cm, 3.4 * cm, 4.0 * cm, 1.6 * cm, 1.4 * cm], font_size=6.8,
        ))
        story.append(Spacer(1, 8))

    # 8 misclassified
    story.append(P("8. Misclassified images", H1))
    any_wrong = False
    for s in summaries:
        for r in s["wrong_rows"]:
            any_wrong = True
            story.append(P(
                f"&bull; <b>{s['split']}</b>: {r['file']} &ndash; true <b>{r['true_class']}</b>, "
                f"predicted <b>{r['predicted_class']}</b> (confidence {float(r['max_confidence']):.4f})<br/>"
                + " ".join(f"{c}={float(r['prob_' + c]):.4f}" for c in CLASSES), BODY))
    if not any_wrong:
        story.append(P("None - every test image was classified correctly.", BODY))

    # 9 conclusion
    story.append(P("9. Conclusion", H1))
    best = max(summaries, key=lambda s: s["accuracy_percent"])
    worst = min(summaries, key=lambda s: s["accuracy_percent"])
    story.append(P(
        f"The trained CNN classified {sum(s['n_correct'] for s in summaries)} of "
        f"{sum(s['n_test_images'] for s in summaries)} test images correctly. "
        f"Best split: <b>{best['split']}</b> with <b>{best['accuracy_percent']:.2f}%</b> accuracy; "
        f"{worst['split']}: <b>{worst['accuracy_percent']:.2f}%</b>.", BODY))
    story.append(Spacer(1, 4))
    story.append(P(
        "The perfect score of the concentrations_based split, which contains every concentration level "
        "(0.01, 0.1, 0.5, 1 and 2 mM), shows that the GAF image representation preserves the information "
        "needed to recognise the substance independently of its concentration. The single error in the "
        "shuffled split is a Benzochinone image at the highest concentration (2 mM) classified as "
        "Idrochinone with 0.60 confidence; the two molecules have very similar quinone / hydroquinone "
        "redox responses, which explains the confusion.", BODY))

    # 10 appendix
    story.append(Spacer(1, 6))
    story.append(P("10. Appendix &ndash; full probability matrix", H1))
    story.append(P("Probability assigned by the CNN to each of the 5 classes, for every test image.", BODY))
    for s in summaries:
        story.append(P(f"{s['split']} split", H2))
        story.append(make_table(
            ["#", "File", "True"] + [c[:11] for c in CLASSES],
            [
                [i, r["file"], r["true_class"]] + [f"{float(r['prob_' + c]):.4f}" for c in CLASSES]
                for i, r in enumerate(s["predictions"])
            ],
            [0.7 * cm, 4.5 * cm, 2.6 * cm, 2.0 * cm, 1.8 * cm, 2.1 * cm, 2.3 * cm, 2.8 * cm], font_size=6.2,
        ))
        story.append(Spacer(1, 10))

    # 11 files
    story.append(P("11. Generated files", H1))
    for path in sorted(os.listdir(results_dir)):
        story.append(P(f"&bull; results/{path}", SMALL))

    doc = SimpleDocTemplate(
        out_path, pagesize=A4,
        leftMargin=1.7 * cm, rightMargin=1.7 * cm, topMargin=1.5 * cm, bottomMargin=1.5 * cm,
        title="GAF-CNN Final Result Report", author="GAF-CNN pipeline",
    )

    def on_page(canvas, doc_):
        canvas.saveState()
        canvas.setFont("DejaVu", 7)
        canvas.setFillColor(colors.HexColor("#5b6570"))
        canvas.drawString(1.7 * cm, 1.0 * cm, "GAF-CNN - Final Result Report")
        canvas.drawRightString(A4[0] - 1.7 * cm, 1.0 * cm, f"Page {doc_.page}")
        canvas.restoreState()

    doc.build(story, onFirstPage=on_page, onLaterPages=on_page)


# --------------------------------------------------------------------------- #

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--results_dir", default="../results")
    ap.add_argument("--data_dir", default="../data/gafimages")
    ap.add_argument("--model", default="CNN_5classes__epo_4000_pat_1000_batch_16_tl.keras")
    args = ap.parse_args()

    now = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S UTC")
    mi = model_info(args.model)
    summaries = collect(args.results_dir, args.data_dir, args.model)
    _, counts = dataset_inventory(args.data_dir)

    txt = os.path.join(args.results_dir, "CNN_Final_Report.txt")
    docx = os.path.join(args.results_dir, "CNN_Final_Report.docx")
    pdf = os.path.join(args.results_dir, "CNN_Final_Report.pdf")

    build_txt(summaries, mi, counts, txt, args.results_dir, now)
    build_docx(summaries, mi, counts, docx, args.results_dir, now)
    build_pdf(summaries, mi, counts, pdf, args.results_dir, now)

    for path in (pdf, docx, txt):
        print(f"  {path}  ({os.path.getsize(path):,} bytes)")


if __name__ == "__main__":
    main()
